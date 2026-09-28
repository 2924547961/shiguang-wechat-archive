"""Render normalized article content to the supported output formats."""

from __future__ import annotations

from email.message import EmailMessage
from email.policy import SMTP
from html import escape
import json
from pathlib import Path

from bs4 import BeautifulSoup


def markdown(content: BeautifulSoup) -> str:
    lines: list[str] = []
    for node in content.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote", "img"]):
        if node.name == "img":
            src = node.get("src") or ""
            if src:
                lines.append(f"![{node.get('alt', '')}]({src})")
        else:
            if node.find_parent(["p", "li", "blockquote"]):
                continue
            text = node.get_text(" ", strip=True)
            if not text:
                continue
            prefix = "#" * int(node.name[1]) + " " if node.name.startswith("h") else "- " if node.name == "li" else "> " if node.name == "blockquote" else ""
            lines.append(prefix + text)
    return "\n\n".join(lines) + "\n"


def html_document(title: str, url: str, author: str, content: BeautifulSoup,
                  extra_html: str = "") -> str:
    metadata = f'<p>公众号：{escape(author)}　来源：<a href="{escape(url, quote=True)}">原文</a></p>'
    return ("<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
            f"<title>{escape(title)}</title>"
            "<style>body{max-width:780px;margin:3rem auto;padding:0 1rem;line-height:1.8}"
            "img{max-width:100%;height:auto}.reply{margin-left:2rem;color:#555}"
            "#js_content,#js_image_content{visibility:visible!important;opacity:1!important;display:block!important}"
            "</style></head><body>"
            f"<h1>{escape(title)}</h1>{metadata}{content}{extra_html}</body></html>")


def export_article(fmt: str, target: Path, title: str, url: str, author: str,
                   content: BeautifulSoup, extra_html: str = "") -> None:
    html = html_document(title, url, author, content, extra_html)
    if fmt == "html":
        target.write_text(html, encoding="utf-8")
    elif fmt == "md":
        extra = BeautifulSoup(extra_html, "html.parser").get_text("\n", strip=True)
        target.write_text(f"# {title}\n\n来源：{url}\n\n" + markdown(content) + "\n" + extra,
                          encoding="utf-8")
    elif fmt == "txt":
        extra = BeautifulSoup(extra_html, "html.parser").get_text("\n", strip=True)
        target.write_text(f"{title}\n{url}\n\n" + content.get_text("\n", strip=True) + "\n" + extra,
                          encoding="utf-8")
    elif fmt == "mhtml":
        _mhtml(target, html, content)
    elif fmt == "docx":
        _docx(target, title, url, author, content, extra_html)
    elif fmt == "pdf":
        _pdf(target, title, url, author, content, extra_html)
    elif fmt == "json":
        target.write_text(json.dumps({"title": title, "url": url, "author": author,
                                      "text": content.get_text("\n", strip=True),
                                      "html": str(content),
                                      "images": [img.get("src", "") for img in content.find_all("img")],
                                      "comments": BeautifulSoup(extra_html, "html.parser").get_text("\n", strip=True)},
                                     ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        raise ValueError(f"不支持的保存格式: {fmt}")


def _mhtml(target: Path, html: str, content: BeautifulSoup) -> None:
    message = EmailMessage(policy=SMTP)
    message["Subject"] = "WeChat article"
    message["MIME-Version"] = "1.0"
    message.set_type("multipart/related")
    root = EmailMessage(policy=SMTP)
    root.set_content(html, subtype="html", charset="utf-8")
    root["Content-Location"] = "article.html"
    message.attach(root)
    for img in content.find_all("img"):
        src = img.get("src") or ""
        image_path = (target.parent / src).resolve()
        if not src.startswith("images/") or not image_path.is_file():
            continue
        ext = image_path.suffix.lower().lstrip(".")
        subtype = "jpeg" if ext in ("jpg", "jpeg") else ext or "octet-stream"
        part = EmailMessage(policy=SMTP)
        part.set_content(image_path.read_bytes(), maintype="image", subtype=subtype)
        part["Content-Location"] = src
        message.attach(part)
    target.write_bytes(message.as_bytes())


def _docx(target: Path, title: str, url: str, author: str,
          content: BeautifulSoup, extra_html: str) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"公众号：{author}    来源：{url}")
    for node in content.find_all(["h1", "h2", "h3", "h4", "p", "li", "img"]):
        if node.name == "img":
            src = node.get("src") or ""
            image_path = target.parent / src
            if src.startswith("images/") and image_path.is_file():
                try:
                    doc.add_picture(str(image_path), width=Inches(5.5))
                except Exception:
                    pass
            continue
        if node.find_parent(["p", "li"]):
            continue
        text = node.get_text(" ", strip=True)
        if not text:
            continue
        if node.name.startswith("h"):
            doc.add_heading(text, min(int(node.name[1]), 4))
        else:
            doc.add_paragraph(text, style="List Bullet" if node.name == "li" else None)
    if extra_html:
        doc.add_heading("评论", 2)
        doc.add_paragraph(BeautifulSoup(extra_html, "html.parser").get_text("\n", strip=True))
    doc.save(target)


def _pdf(target: Path, title: str, url: str, author: str,
         content: BeautifulSoup, extra_html: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image, SimpleDocTemplate, Paragraph, Spacer

    font = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font))
    normal = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17,
                            textColor=colors.black, alignment=TA_LEFT, spaceAfter=9)
    heading = ParagraphStyle("heading", parent=normal, fontSize=17, leading=24, spaceAfter=15)
    story = [Paragraph(escape(title), heading),
             Paragraph(escape(f"公众号：{author}    来源：{url}"), normal), Spacer(1, 10)]
    for node in content.find_all(["h1", "h2", "h3", "h4", "p", "li", "img"]):
        if node.name == "img":
            src = node.get("src") or ""
            image_path = target.parent / src
            if src.startswith("images/") and image_path.is_file():
                try:
                    figure = Image(str(image_path))
                    scale = min(1.0, 480 / figure.imageWidth, 620 / figure.imageHeight)
                    figure.drawWidth = figure.imageWidth * scale
                    figure.drawHeight = figure.imageHeight * scale
                    story.extend([figure, Spacer(1, 10)])
                except Exception:
                    pass
            continue
        if node.find_parent(["p", "li"]):
            continue
        text = node.get_text(" ", strip=True)
        if text:
            story.append(Paragraph(escape(text), heading if node.name.startswith("h") else normal))
    if extra_html:
        story.append(Paragraph("评论", heading))
        story.append(Paragraph(escape(BeautifulSoup(extra_html, "html.parser").get_text(" ", strip=True)), normal))
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=50, rightMargin=50).build(story)
