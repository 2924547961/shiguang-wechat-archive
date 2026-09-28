from __future__ import annotations

import csv
import datetime as dt
import html
import json
import os
import shutil
import uuid
from pathlib import Path

from .common import STATIC, Cancelled, check_cancel, inside, safe_name
from .messages import KIND_NAMES
from .store import Archive, create_store, finalize, set_meta, unpack_message, friends_clause


def csv_cell(value):
    value = "" if value is None else str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def readable(message):
    detail = message["detail"]
    if isinstance(detail, str):
        detail = json.loads(detail or "{}")
    body = message.get("body") or ""
    parts = [body]
    if detail.get("description") and detail["description"] != body:
        parts.append(detail["description"])
    if detail.get("url"):
        parts.append(detail["url"])
    if detail.get("quote"):
        q = detail["quote"]
        parts.append("引用 " + q.get("sender", "") + ": " + q.get("body", ""))
    def forward(items, depth=1):
        for item in items:
            parts.append("  " * depth + (item.get("sender_name", "") + ": " if item.get("sender_name") else "") + (item.get("body") or KIND_NAMES.get(item.get("kind"), "消息")))
            forward(item.get("items", []), depth + 1)
    forward(detail.get("items", []))
    if message.get("media_status") and not message.get("media_path"):
        parts.append("[" + message["media_status"] + "]")
    return "\n".join(parts)


def copy_media(message, archive_root, output, enabled=True, cancel=None):
    def copy_one(item):
        rel = item.get("media_path", "")
        if not enabled:
            item["media_path"] = ""
            if rel:
                item["media_status"] = "此次导出未包含媒体"
            return
        if not rel:
            return
        try:
            source = inside(archive_root, Path(archive_root) / rel)
            if item.get("kind")=="video" and source.is_file():
                from .playback import compatible_video
                source=compatible_video(source,cancel)
                rel=source.relative_to(archive_root).as_posix()
                item["media_path"]=rel
            dest = inside(output, output / rel)
            if source.is_file():
                dest.parent.mkdir(parents=True, exist_ok=True)
                if not dest.exists():
                    shutil.copy2(source, dest)
            else:
                item["media_path"], item["media_status"] = "", "归档中的媒体文件已被移走"
        except (ValueError, OSError):
            item["media_path"], item["media_status"] = "", "媒体复制失败"
    copy_one(message)
    avatar = {'media_path': message.get('sender_avatar', '')}
    copy_one(avatar)
    message['sender_avatar'] = avatar.get('media_path', '')
    def walk(items):
        for item in items:
            copy_one(item); walk(item.get("items", []))
    walk(message.get("detail", {}).get("items", []))


def html_document(path, title, messages):
    shell = (STATIC / "archive.html").read_text("utf-8")
    style = (STATIC / "reader.css").read_text("utf-8")
    renderer = (STATIC / "renderer.js").read_text("utf-8")
    js = (STATIC / "archive.js").read_text("utf-8")
    data = json.dumps({"title": title, "messages": messages}, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    page = shell.replace("{{TITLE}}", html.escape(title)).replace("{{STYLE}}", style).replace("{{DATA}}", data).replace("{{RENDERER}}", renderer).replace("{{SCRIPT}}", js)
    path.write_text(page, "utf-8")


def word_document(path, title, messages, output, cancel=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.8)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(7)
    document.add_heading(title, 0)
    document.add_paragraph(f"拾光 · 本地聊天归档  |  共 {len(messages):,} 条消息")
    document.add_paragraph("导出时间：" + dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    def link(paragraph, text, url):
        relation = document.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        anchor = OxmlElement("w:hyperlink"); anchor.set(qn("r:id"), relation)
        run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "24725B"); props.append(color); run.append(props)
        content = OxmlElement("w:t"); content.text = text; run.append(content); anchor.append(run); paragraph._p.append(anchor)
    for index, m in enumerate(messages):
        if index % 50 == 0:
            check_cancel(cancel)
        meta = document.add_paragraph()
        meta.paragraph_format.keep_with_next = True
        run = meta.add_run(m["time"] + "  " + (m.get("sender_name") or "未知发送者") + " · " + KIND_NAMES.get(m["kind"], "消息"))
        run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string("24725B")
        body = readable(m)
        if body:
            document.add_paragraph(body)
        rel = m.get("media_path")
        if rel:
            media = inside(output, output / rel)
            if m["kind"] in {"image", "emoji"} and media.suffix.lower() in {".jpg", ".jpeg", ".png", ".bmp", ".gif"}:
                try:
                    document.add_picture(str(media), width=Inches(3.6))
                except Exception:
                    link(document.add_paragraph(), "打开图片文件", rel)
            else:
                link(document.add_paragraph(), "打开" + KIND_NAMES.get(m["kind"], "附件"), rel)
    footer = section.footer.paragraphs[0]
    footer.text = "拾光 SHIGUANG  ·  保存在此刻，重逢于未来"
    footer.runs[0].font.size = Pt(8)
    document.save(path)


def export_contacts(archive: Archive, path: Path, selected=None):
    with archive.connect() as c, path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f); writer.writerow(["微信ID", "昵称", "备注", "微信号", "类型"])
        for r in c.execute("SELECT username,nickname,remark,alias,kind FROM contacts WHERE " + friends_clause(c) + " ORDER BY remark,nickname"):
            if selected is None or r[0] in selected:
                writer.writerow([csv_cell(v) for v in r])


def export_data(archive: Archive, options, output_root, progress, cancel=None):
    formats = list(dict.fromkeys(options.get("formats", ["html"])))
    if not formats or any(f not in {"html", "txt", "csv", "docx", "sqlite", "contacts", "emoji"} for f in formats):
        raise ValueError("请选择有效的导出格式。")
    cids = options.get("conversations")
    if cids is not None and not isinstance(cids, list):
        raise ValueError("会话选择无效。")
    start, end = options.get("start"), options.get("end")
    if options.get("include_media", True) and not set(formats)<= {"contacts","emoji"} and "html" not in formats:
        formats = [*formats, "html"]
    if start and end and int(end) <= int(start):
        raise ValueError("结束日期需要晚于开始日期。")
    output_root = Path(output_root).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    name = "拾光导出_" + dt.datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6]
    output = output_root / ("." + name + ".partial")
    output.mkdir()
    dest = None
    manifest, count = [], 0
    media_summary = {'available':0,'missing':0,'by_type':{}}
    missing_media = []
    try:
        with archive.connect() as c:
            chats = [dict(r) for r in c.execute("SELECT * FROM conversations ORDER BY last_ts DESC")]
        if cids is not None:
            chats = [r for r in chats if r["id"] in cids]
        if formats == ["contacts"]:
            chats = []
        if not chats and any(f != "contacts" for f in formats):
            raise ValueError("没有选中可导出的会话。")
        if "contacts" in formats:
            export_contacts(archive, output / "联系人.csv", options.get("contacts"))
            manifest.append({"file": "联系人.csv", "format": "contacts"})
        if "sqlite" in formats:
            dest = create_store(output / "聊天记录.sqlite")
            set_meta(dest, "exported_at", dt.datetime.now().isoformat())
            set_meta(dest, "account", archive.metadata().get("account"))
            set_meta(dest, "source", "filtered_export")
            with archive.connect() as src:
                for row in src.execute("SELECT * FROM contacts"):
                    contact = dict(row)
                    avatar = {"media_path": contact.get("avatar", "")}
                    copy_media(avatar, archive.directory, output, options.get("include_media", True))
                    contact["avatar"] = avatar.get("media_path", "")
                    dest.execute("INSERT INTO contacts VALUES(?,?,?,?,?,?)", tuple(contact[k] for k in ['username', 'nickname', 'remark', 'alias', 'kind', 'avatar']))
        for i, chat in enumerate(chats):
            check_cancel(cancel)
            progress("export", int(95 * i / max(1, len(chats))), f"正在导出会话 · {i + 1} / {len(chats)}")
            messages = []
            for index, row in enumerate(archive.iterate([chat["id"]], start, end)):
                if index % 100 == 0:
                    check_cancel(cancel)
                original = dict(row)
                original.pop('sender_avatar', None)
                msg = unpack_message(row)
                if formats==['emoji'] and msg['kind']!='emoji':
                    continue
                copy_media(msg, archive.directory, output,
                           options.get("include_media", True) or ('emoji' in formats and msg['kind']=='emoji'), cancel)
                def media_inventory(item):
                    kind=item.get('kind')
                    if kind in {'image','video','audio','emoji','file'}:
                        status='available' if item.get('media_path') else 'missing'
                        media_summary[status]+=1
                        media_summary['by_type'].setdefault(kind,{'available':0,'missing':0})[status]+=1
                        if status=='missing':missing_media.append({'conversation':chat['title'],'message_id':msg.get('server_id'),'kind':kind,'reason':item.get('media_status') or '未找到本地媒体'})
                    for child in item.get('detail',{}).get('items',item.get('items',[])):media_inventory(child)
                media_inventory(msg)
                if dest is not None:
                    original["media_path"], original["media_status"] = msg.get("media_path", ""), msg.get("media_status", "")
                    original["detail"] = json.dumps(msg["detail"], ensure_ascii=False)
                    cols = list(original)
                    dest.execute("INSERT INTO messages(" + ",".join(cols) + ") VALUES(" + ",".join("?" for _ in cols) + ")", [original[k] for k in cols])
                messages.append(msg)
            if not messages:
                continue
            count += len(messages)
            basename = safe_name(chat["title"]) + "_" + __import__("hashlib").sha256(chat["id"].encode()).hexdigest()[:8]
            if dest is not None:
                dest.execute("INSERT INTO conversations(id,title,kind) VALUES(?,?,?)", (chat["id"], chat["title"], chat["kind"]))
            for fmt in formats:
                if fmt == 'emoji':
                    folder=output/('表情包_'+basename)
                    folder.mkdir(exist_ok=True)
                    cards=[];seen=set()
                    for m in messages:
                        if m['kind']!='emoji' or not m.get('media_path'):continue
                        source=inside(output,output/m['media_path'])
                        if not source.is_file():continue
                        fingerprint=__import__('hashlib').sha256(source.read_bytes()).hexdigest()
                        if fingerprint in seen:continue
                        seen.add(fingerprint)
                        filename=fingerprint[:16]+source.suffix.lower()
                        shutil.copy2(source,folder/filename)
                        label=html.escape(m.get('time','')+' · '+(m.get('media_status') or '本机文件'))
                        href=html.escape(filename,quote=True)
                        cards.append(f'<a href="{href}" target="_blank"><img src="{href}" loading="lazy" alt="表情包"><span>{label}</span></a>')
                    (folder/'index.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
                        '<title>表情包导出</title><style>body{font:14px sans-serif;background:#f6f8f4;color:#274234;padding:24px}'
                        'main{display:grid;grid-template-columns:repeat(auto-fill,minmax(160px,1fr));gap:12px}'
                        'a{background:white;padding:12px;text-align:center;border-radius:8px;color:inherit;text-decoration:none}'
                        'img{width:100%;height:150px;object-fit:contain}span{display:block;margin-top:8px}</style>'
                        f'<h1>{html.escape(chat["title"])} · 表情包</h1><p>可打开或另存本机已恢复文件，共 {len(cards)} 个；缩略图、首帧会按状态标注。</p>'
                        '<main>'+''.join(cards)+'</main></html>','utf-8')
                    manifest.append({'title':chat['title']+' · 表情包','file':folder.name+'/index.html','format':'emoji','messages':len(cards)})
                    continue
                path = output / (basename + "." + fmt)
                if fmt == "html":
                    html_document(path, chat["title"], messages)
                elif fmt == "txt":
                    with path.open("w", encoding="utf-8") as f:
                        f.write(chat["title"] + "\n" + "=" * 50 + "\n\n")
                        for m in messages:
                            f.write(f"[{m['time']}] {m.get('sender_name', '')} [{KIND_NAMES.get(m['kind'], '消息')}]\n{readable(m)}\n")
                            if m.get("media_path"):
                                f.write("附件: " + m["media_path"] + "\n")
                            f.write("\n")
                elif fmt == "csv":
                    with path.open("w", encoding="utf-8-sig", newline="") as f:
                        writer = csv.writer(f)
                        writer.writerow(["时间", "发送者", "发送者ID", "是否本人", "类型", "内容", "消息ID", "附件", "媒体状态"])
                        for m in messages:
                            writer.writerow([csv_cell(x) for x in [m["time"], m.get("sender_name"), m.get("sender"), m.get("is_self"), KIND_NAMES.get(m["kind"], "消息"), readable(m), m.get("server_id"), m.get("media_path"), m.get("media_status")]])
                elif fmt == "docx":
                    word_document(path, chat["title"], messages, output, cancel)
                else:
                    continue
                manifest.append({"title": chat["title"], "file": path.name, "format": fmt, "messages": len(messages)})
        if dest is not None:
            finalize(dest, lambda *a: None); dest.close(); dest = None
            manifest.append({"file": "聊天记录.sqlite", "format": "sqlite", "messages": count})
        if count == 0 and any(f != "contacts" for f in formats):
            raise ValueError("所选时间范围内没有消息。")
        (output / "导出清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), "utf-8")
        (output / '媒体恢复情况.json').write_text(json.dumps({'summary':media_summary,'missing':missing_media},ensure_ascii=False,indent=2),'utf-8')
        links = "".join('<li><a href="' + html.escape(r["file"], quote=True) + '">' + html.escape(r.get("title", r["file"])) + '</a><span>' + r["format"].upper() + '</span></li>' for r in manifest)
        (output / "打开归档.html").write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>拾光 · 导出归档</title><style>body{font:16px/1.8 "Microsoft YaHei",sans-serif;background:#f4f7f4;color:#223c33;max-width:900px;margin:50px auto;padding:24px}h1{font-size:36px}li{background:white;list-style:none;margin:10px 0;padding:15px 22px;border-radius:12px;display:flex;justify-content:space-between}ul{padding:0}a{color:#28775c;text-decoration:none}span{color:#7c9187;font-size:12px}</style><h1>这一刻，已被珍藏。</h1><p>拾光 · 本地聊天归档</p><ul>' + links + '</ul></html>', "utf-8")
        check_cancel(cancel)
        final = output_root / name
        os.replace(output, final)
        progress("done", 100, "导出完成")
        return {"path": str(final), "index": str(final / "打开归档.html"), "messages": count, "files": len(manifest), 'media':media_summary}
    except BaseException:
        if dest is not None:
            dest.close()
        shutil.rmtree(inside(output_root, output), ignore_errors=True)
        raise


def export_report(archive, year, cid, output_root, progress, cancel=None):
    progress("report", 10, "正在整理这一年的对话")
    data = archive.annual(year, cid, cancel)
    if not data["total"]:
        raise ValueError("这一年没有聊天记录，请换一个年份。")
    progress("report", 85, "正在制作年度回顾")
    output_root = Path(output_root).expanduser().resolve(); output_root.mkdir(parents=True, exist_ok=True)
    folder = output_root / ("年度回顾_" + str(year) + "_" + uuid.uuid4().hex[:8]); folder.mkdir()
    try:
        shell = (STATIC / "report.html").read_text("utf-8")
        encoded = json.dumps(data, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
        page = shell.replace("{{DATA}}", encoded)
        path = folder / "年度回顾.html"
        path.write_text(page, "utf-8")
        from .poster import annual_poster
        image = folder / "年度回顾.png"
        annual_poster(data, image)
        check_cancel(cancel)
        progress("done", 100, "年度回顾与高清图片已生成")
        return {"path": str(folder), "index": str(path), "image": str(image), "report": data}
    except BaseException:
        shutil.rmtree(inside(output_root, folder), ignore_errors=True)
        raise
